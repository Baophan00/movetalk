#!/usr/bin/env python3
"""Extract content from PPTX files and generate lesson data for L02-L07."""
import json
import os
import re
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
PPTX_DIR = os.path.expanduser("~/.hermes/cache/documents")

# Map lesson number to PPTX file
PPTX_FILES = {
    2: "doc_85b71c6a3223_MOVE_ACTIVATE_1_L02_EXPAND_THE_SUBJECT_v2 (1).pptx",
    3: "doc_31bbfb35b0ed_MOVE_ACTIVATE_1_L03_BUILD_THE_MESSAGE.pptx",
    4: "doc_373fb946cf8f_MOVE_ACTIVATE_1_L04_BUILD_THE_NOUN_PHRASE_V3.pptx",
    5: "doc_60397f89e651_MOVE_ACTIVATE_1_L05_POINT_ASK_CONNECT.pptx",
    6: "doc_0601252e0ada_MOVE_ACTIVATE_1_L06_THERE_IS_THERE_ARE.pptx",
    7: "doc_c7f9425e500d_MOVE_ACTIVATE_1_L07_CONFUSING_QUANTIFIERS_DETERMINERS .pptx",
}

def extract_slides(filepath):
    """Extract all text from slides."""
    prs = Presentation(filepath)
    slides = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            slides.append(texts)
    return slides

def extract_vocab_from_slides(slides):
    """Extract vocabulary words from slides."""
    vocab = []
    for slide in slides:
        text = ' '.join(slide).lower()
        # Look for vocabulary patterns
        if 'word' in text or 'vocabulary' in text or 'từ vựng' in text:
            for line in slide:
                # Simple extraction: look for English words with Vietnamese
                if '·' in line or '—' in line or '(' in line:
                    vocab.append(line)
    return vocab[:20]  # Limit to 20 words

def extract_conversations_from_slides(slides):
    """Extract conversation dialogs from slides."""
    convs = []
    for i, slide in enumerate(slides):
        text = ' '.join(slide).lower()
        if 'dialog' in text or 'hội thoại' in text or 'conversation' in text:
            # Try to extract speaker lines
            lines = []
            for line in slide:
                if ':' in line or '·' in line:
                    lines.append(line)
            if len(lines) >= 4:
                convs.append({
                    'title': slide[0] if slide else f'Conversation {i}',
                    'lines': lines[:10]
                })
    return convs[:5]  # Limit to 5 conversations

def generate_lesson_data(lesson_num, slides):
    """Generate lesson data structure."""
    # Extract title from first slide
    title = "Unknown"
    if slides and slides[0]:
        for line in slides[0]:
            if 'L0' in line and '·' in line:
                title = line.split('·')[-1].strip()
                break
    
    # Generate basic structure
    lesson = {
        "title": title,
        "vi": f"Bài {lesson_num}",
        "goal": f"By the end of this lesson, you can master the key concepts of {title}.",
        "today": ["Review previous lesson", "Learn new patterns", "Practice with partner", "Build confidence"],
        "parts": {
            "vocab": [
                {"t": "h", "text": "Vocabulary"},
                {"t": "flashcard", "items": []}
            ],
            "grammar": [
                {"t": "h", "text": "Grammar"},
                {"t": "table", "head": ["Pattern", "Example", "Use"], "rows": []}
            ],
            "reading": [
                {"t": "h", "text": "Reading"},
                {"t": "p", "text": "Reading content will be added."}
            ],
            "conversations": [],
            "listening": [
                {"t": "h", "text": "Listening"},
                {"t": "say", "items": []}
            ],
            "exercises": [
                {"t": "h", "text": "Exercises"},
                {"t": "quiz", "items": []}
            ],
            "speaking": [
                {"t": "h", "text": "Speaking"},
                {"t": "p", "text": "Practice speaking with a partner."}
            ],
            "roleplay": [
                {"t": "steps", "title": "Role-play", "context": "Practice with a partner.", "steps": ["Step 1", "Step 2", "Step 3"], "useful": "Useful phrases"}
            ],
            "finalTask": [
                {"t": "label", "text": "Final Task"},
                {"t": "p", "text": "Complete the final task."}
            ],
            "reference": [
                {"t": "h", "text": "Reference"},
                {"t": "note", "text": "Review key points."}
            ]
        }
    }
    
    # Extract vocab
    vocab_items = extract_vocab_from_slides(slides)
    if vocab_items:
        lesson["parts"]["vocab"][1]["items"] = [
            {"word": v.split('·')[0].strip() if '·' in v else v, "phonetic": "", "vi": "", "def": "", "examples": []}
            for v in vocab_items[:10]
        ]
    
    # Extract conversations
    convs = extract_conversations_from_slides(slides)
    if convs:
        lesson["parts"]["conversations"] = [
            {"t": "dialog", "title": c["title"], "lines": [{"speaker": "A", "text": l} for l in c["lines"]]}
            for c in convs
        ]
    
    return lesson

def main():
    for lesson_num, filename in PPTX_FILES.items():
        filepath = os.path.join(PPTX_DIR, filename)
        if not os.path.exists(filepath):
            print(f"Warning: {filename} not found")
            continue
        
        print(f"Processing L{lesson_num}: {filename}")
        slides = extract_slides(filepath)
        lesson = generate_lesson_data(lesson_num, slides)
        
        out = os.path.join(DATA, f"L{lesson_num:02d}.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(lesson, f, ensure_ascii=False, indent=2)
        print(f"  -> {out}")

if __name__ == "__main__":
    main()
