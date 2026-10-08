#!/usr/bin/env python3
"""Extract content from PPTX files and generate lesson data for L02-L07.
Smart extraction: analyze slide content to identify vocab, grammar, conversations.
"""
import json
import os
import re
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
PPTX_DIR = os.path.expanduser("~/.hermes/cache/documents")

PPTX_FILES = {
    2: "doc_85b71c6a3223_MOVE_ACTIVATE_1_L02_EXPAND_THE_SUBJECT_v2 (1).pptx",
    3: "doc_31bbfb35b0ed_MOVE_ACTIVATE_1_L03_BUILD_THE_MESSAGE.pptx",
    4: "doc_373fb946cf8f_MOVE_ACTIVATE_1_L04_BUILD_THE_NOUN_PHRASE_V3.pptx",
    5: "doc_60397f89e651_MOVE_ACTIVATE_1_L05_POINT_ASK_CONNECT.pptx",
    6: "doc_0601252e0ada_MOVE_ACTIVATE_1_L06_THERE_IS_THERE_ARE.pptx",
    7: "doc_c7f9425e500d_MOVE_ACTIVATE_1_L07_CONFUSING_QUANTIFIERS_DETERMINERS .pptx",
}

def extract_slides(filepath):
    prs = Presentation(filepath)
    slides = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            slides.append(texts)
    return slides

def is_vocab_slide(slide):
    """Check if slide contains vocabulary."""
    text = ' '.join(slide).lower()
    return any(kw in text for kw in ['word', 'vocabulary', 'từ vựng', 'flashcard', 'từ mới'])

def is_conversation_slide(slide):
    """Check if slide contains conversation."""
    text = ' '.join(slide).lower()
    return any(kw in text for kw in ['dialog', 'hội thoại', 'conversation', 'mẫu hội thoại'])

def is_grammar_slide(slide):
    """Check if slide contains grammar."""
    text = ' '.join(slide).lower()
    return any(kw in text for kw in ['grammar', 'ngữ pháp', 'pattern', 'công thức', 'structure'])

def extract_vocab_from_slide(slide):
    """Extract vocabulary items from a slide."""
    items = []
    for line in slide:
        # Skip headers and footers
        if '©' in line or 'MOVE' in line or 'Phoebe' in line:
            continue
        # Look for patterns like "word · meaning" or "word — meaning"
        if '·' in line or '—' in line:
            parts = re.split(r'[·—]', line, maxsplit=1)
            if len(parts) == 2:
                word = parts[0].strip()
                meaning = parts[1].strip()
                # Clean up word
                word = re.sub(r'^[\d\.\s]+', '', word)
                if word and len(word) < 50:
                    items.append({
                        "word": word,
                        "phonetic": "",
                        "vi": meaning,
                        "def": "",
                        "examples": []
                    })
    return items

def extract_conversation_from_slide(slide):
    """Extract conversation from a slide."""
    lines = []
    for line in slide:
        if '©' in line or 'MOVE' in line or 'Phoebe' in line:
            continue
        # Look for speaker patterns
        if ':' in line or '·' in line:
            parts = re.split(r'[:·]', line, maxsplit=1)
            if len(parts) == 2:
                speaker = parts[0].strip()
                text = parts[1].strip()
                if speaker and text and len(speaker) < 20:
                    lines.append({"speaker": speaker, "text": text})
    return lines

def generate_lesson_data(lesson_num, slides):
    """Generate lesson data structure."""
    # Extract title
    title = f"Lesson {lesson_num}"
    if slides and slides[0]:
        for line in slides[0]:
            if 'L0' in line and '·' in line:
                title = line.split('·')[-1].strip()
                break
    
    # Extract content
    vocab_items = []
    conversations = []
    grammar_rows = []
    
    for slide in slides:
        if is_vocab_slide(slide):
            vocab_items.extend(extract_vocab_from_slide(slide))
        elif is_conversation_slide(slide):
            conv_lines = extract_conversation_from_slide(slide)
            if len(conv_lines) >= 4:
                conversations.append({
                    "title": slide[0] if slide else f"Conversation {len(conversations)+1}",
                    "lines": conv_lines
                })
        elif is_grammar_slide(slide):
            # Extract grammar patterns
            for line in slide:
                if '·' in line or '—' in line:
                    parts = re.split(r'[·—]', line, maxsplit=2)
                    if len(parts) >= 2:
                        grammar_rows.append([p.strip() for p in parts[:3]])
    
    # Build lesson
    lesson = {
        "title": title,
        "vi": f"Bài {lesson_num}",
        "goal": f"By the end of this lesson, you can master the key concepts of {title}.",
        "today": ["Review previous lesson", "Learn new patterns", "Practice with partner", "Build confidence"],
        "parts": {
            "vocab": [
                {"t": "h", "text": "Vocabulary"},
                {"t": "flashcard", "items": vocab_items[:15] if vocab_items else []}
            ],
            "grammar": [
                {"t": "h", "text": "Grammar"},
                {"t": "table", "head": ["Pattern", "Example", "Use"], "rows": grammar_rows[:10] if grammar_rows else []}
            ],
            "reading": [
                {"t": "h", "text": "Reading"},
                {"t": "p", "text": "Reading content will be added."}
            ],
            "conversations": [
                {"t": "dialog", "title": c["title"], "lines": c["lines"]}
                for c in conversations[:5]
            ] if conversations else [],
            "listening": [
                {"t": "h", "text": "Listening"},
                {"t": "say", "items": []}
            ],
            "exercises": [
                {"t": "h", "text": "Exercises"},
                {"t": "quiz", "items": []}
            ],
            "speaking": [
                {"t": "h", "text": "Speaking"},
                {"t": "p", "text": "Practice speaking with a partner."}
            ],
            "roleplay": [
                {"t": "steps", "title": "Role-play", "context": "Practice with a partner.", "steps": ["Step 1", "Step 2", "Step 3"], "useful": "Useful phrases"}
            ],
            "finalTask": [
                {"t": "label", "text": "Final Task"},
                {"t": "p", "text": "Complete the final task."}
            ],
            "reference": [
                {"t": "h", "text": "Reference"},
                {"t": "note", "text": "Review key points."}
            ]
        }
    }
    
    return lesson

def main():
    for lesson_num, filename in PPTX_FILES.items():
        filepath = os.path.join(PPTX_DIR, filename)
        if not os.path.exists(filepath):
            print(f"Warning: {filename} not found")
            continue
        
        print(f"Processing L{lesson_num}: {filename}")
        slides = extract_slides(filepath)
        lesson = generate_lesson_data(lesson_num, slides)
        
        out = os.path.join(DATA, f"L{lesson_num:02d}.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(lesson, f, ensure_ascii=False, indent=2)
        print(f"  -> {out}")
        print(f"     vocab: {len(lesson['parts']['vocab'][1]['items'])}, convs: {len(lesson['parts']['conversations'])}")

if __name__ == "__main__":
    main()
